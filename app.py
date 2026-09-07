import os
import io
import re
import copy
import base64
import unicodedata
import pandas as pd
import streamlit as st
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Font, Border, Side
from deep_translator import GoogleTranslator
from github import Github

st.set_page_config(page_title="Excel Mirror Translator", page_icon="📊", layout="wide")

DEFAULT_DICT_PATH = "dictionar.xlsx"

BORDER_MISSING = Border(
    left=Side(style='medium', color='FFC000'),
    right=Side(style='medium', color='FFC000'),
    top=Side(style='medium', color='FFC000'),
    bottom=Side(style='medium', color='FFC000')
)

BORDER_NOT_FOUND = Border(
    left=Side(style='medium', color='FF0000'),
    right=Side(style='medium', color='FF0000'),
    top=Side(style='medium', color='FF0000'),
    bottom=Side(style='medium', color='FF0000')
)

def normalize_text(text: str) -> str:
    if not isinstance(text, str):
        text = str(text)
    text = text.lower().strip()
    replacements = {'ă': 'a', 'â': 'a', 'î': 'i', 'ș': 's', 'ț': 't', 'ş': 's', 'ţ': 't'}
    for k, v in replacements.items():
        text = text.replace(k, v)
    nfd = unicodedata.normalize('NFD', text)
    text = ''.join(c for c in nfd if unicodedata.category(c) != 'Mn')
    text = re.sub(r'\s+', ' ', text)
    return text.strip()

def is_purely_numeric(val) -> bool:
    if val is None:
        return False
    if isinstance(val, (int, float)):
        return True
    s = str(val).strip()
    has_digit = any(c.isdigit() for c in s)
    only_chars = bool(re.match(r'^[\s\d.,\-–%/:]+$', s))
    return (has_digit and only_chars) or (s == "-")

def translate_with_google(text: str) -> tuple[str, str]:
    try:
        translated = GoogleTranslator(source='ro', target='en').translate(text)
        if not translated:
            return text, "NOT FOUND"
        if translated.lower() == text.lower() and len(text) <= 4:
            return text, "NOT FOUND"
        return translated, "MISSING"
    except Exception:
        return text, "NOT FOUND"

def extract_pairs_from_sheet(sheet):
    pairs = {}
    if sheet.max_row is None or sheet.max_row < 1:
        return pairs

    src_col, trg_col = 1, 2
    header_found = False

    for r in range(1, min(15, sheet.max_row + 1)):
        row_vals = [normalize_text(str(sheet.cell(r, c).value or '')) for c in range(1, min(15, sheet.max_column + 1))]
        for idx, val in enumerate(row_vals):
            if any(k in val for k in ["ro", "sursa", "romana", "original", "source"]):
                src_col = idx + 1
                header_found = True
            if any(k in val for k in ["en", "tinta", "engleza", "traducere", "target", "translation"]):
                trg_col = idx + 1
                header_found = True

    start_row = 2 if header_found else 1
    if src_col == trg_col:
        src_col, trg_col = 1, 2

    for r in range(start_row, sheet.max_row + 1):
        src = sheet.cell(r, src_col).value
        trg = sheet.cell(r, trg_col).value

        if not src or not trg:
            non_empty = [c for c in range(1, min(10, sheet.max_column + 1)) if sheet.cell(r, c).value not in [None, ""]]
            if len(non_empty) >= 2:
                src = sheet.cell(r, non_empty[0]).value
                trg = sheet.cell(r, non_empty[1]).value

        if src and trg:
            src_str = str(src).strip()
            trg_str = str(trg).strip()
            if src_str and trg_str and not is_purely_numeric(src_str):
                pairs[normalize_text(src_str)] = trg_str

    return pairs

def load_master_dictionary():
    dictionary = {}
    if os.path.exists(DEFAULT_DICT_PATH):
        try:
            dict_wb = load_workbook(DEFAULT_DICT_PATH, data_only=True)
            for sheet_name in dict_wb.sheetnames:
                sheet = dict_wb[sheet_name]
                dictionary.update(extract_pairs_from_sheet(sheet))
        except Exception:
            pass
    return dictionary

def update_github_dictionary(new_approved_pairs: dict):
    token = st.secrets.get("GITHUB_TOKEN")
    repo_name = st.secrets.get("GITHUB_REPO", "stefanbogdanemail-dev/excel-mirror-translator")
    
    if not token:
        st.warning("GITHUB_TOKEN nu este setat în Secrets. Termenii nu au putut fi sincronizați automat pe GitHub.")
        return False

    try:
        g = Github(token)
        repo = g.get_repo(repo_name)
        contents = repo.get_contents(DEFAULT_DICT_PATH)
        file_bytes = base64.b64decode(contents.content)
        
        dict_wb = load_workbook(io.BytesIO(file_bytes))
        sheet = dict_wb.active

        for src, trg in new_approved_pairs.items():
            sheet.append([src, trg])

        out_buf = io.BytesIO()
        dict_wb.save(out_buf)
        out_buf.seek(0)

        repo.update_file(
            path=DEFAULT_DICT_PATH,
            message="Auto-update dictionar din aplicatia Streamlit (termeni aprobati)",
            content=out_buf.getvalue(),
            sha=contents.sha
        )
        return True
    except Exception as e:
        st.error(f"Eroare la actualizarea pe GitHub: {e}")
        return False

def run_translation_pipeline(input_bytes, master_dict):
    wb = load_workbook(io.BytesIO(input_bytes), data_only=False)
    wb_vals = load_workbook(io.BytesIO(input_bytes), data_only=True)

    dictionary = dict(master_dict)
    for s in wb.sheetnames:
        if normalize_text(s) in ["dictionar", "dictionary", "glosar"]:
            dictionary.update(extract_pairs_from_sheet(wb[s]))

    review_rows = []
    stats = {"MATCH": 0, "MISSING": 0, "NOT FOUND": 0, "SKIP": 0}
    unique_new_terms = {}

    for sheet_name in wb.sheetnames:
        src_sheet = wb[sheet_name]
        val_sheet = wb_vals[sheet_name]
        s_norm = normalize_text(sheet_name)

        if s_norm in ["dictionar", "dictionary", "glosar"] or sheet_name.startswith('_'):
            continue

        tr_twin = wb.copy_worksheet(src_sheet)
        tr_title = f"{sheet_name[:25]} (EN)"[:31]
        tr_twin.title = tr_title

        target_idx = wb.index(src_sheet) + 1
        current_idx = wb.index(tr_twin)
        wb.move_sheet(tr_twin, offset=target_idx - current_idx)

        merged_non_anchors = set()
        for rng in tr_twin.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if not (r == rng.min_row and c == rng.min_col):
                        merged_non_anchors.add((r, c))

        for r in range(1, tr_twin.max_row + 1):
            for c in range(1, tr_twin.max_column + 1):
                if (r, c) in merged_non_anchors:
                    stats["SKIP"] += 1
                    continue

                dst_c = tr_twin.cell(r, c)
                val = val_sheet.cell(r, c).value
                
                if dst_c.data_type == 'f':
                    dst_c.value = val

                if val is None or str(val).strip() == "" or is_purely_numeric(val):
                    stats["SKIP"] += 1
                    continue

                val_str = str(val).strip()
                val_norm = normalize_text(val_str)

                if val_norm in dictionary:
                    dst_c.value = dictionary[val_norm]
                    stats["MATCH"] += 1
                else:
                    tr_text, status = translate_with_google(val_str)
                    dst_c.value = tr_text
                    stats[status] += 1
                    if status == "MISSING":
                        dst_c.border = BORDER_MISSING
                        review_rows.append((tr_title, r, c, val_str, tr_text, "MISSING"))
                        if val_str not in unique_new_terms:
                            unique_new_terms[val_str] = tr_text
                    else:
                        dst_c.border = BORDER_NOT_FOUND
                        review_rows.append((tr_title, r, c, val_str, tr_text, "NOT FOUND"))

    summary = wb.create_sheet(title="_Review Summary", index=0)
    summary.append(["Sheet", "Row", "Column", "Original", "Translation", "Status"])
    for col_idx in range(1, 7):
        c = summary.cell(1, col_idx)
        c.font = Font(bold=True)
        c.fill = PatternFill(start_color="D9D9D9", fill_type="solid")

    for row_data in review_rows:
        summary.append(list(row_data))
        status_cell = summary.cell(summary.max_row, 6)
        color = "FF0000" if row_data[5] == "NOT FOUND" else "FFC000"
        status_cell.fill = PatternFill(start_color=color, fill_type="solid")

    return wb, stats, unique_new_terms

# --- UI Streamlit ---
st.title("📊 Excel Mirror Translator (Auto-Sync Dicționar)")

master_dict = load_master_dictionary()
if master_dict:
    st.success(f"Dicționarul master este activ ({len(master_dict)} termeni unici).")
else:
    st.warning("Fișierul `dictionar.xlsx` nu a fost găsit.")

uploaded_file = st.file_uploader("Încarcă fișierul Excel de Tradus (.xlsx)", type=["xlsx"])

if "translated_wb" not in st.session_state:
    st.session_state.translated_wb = None
    st.session_state.new_terms_df = None
    st.session_state.stats = None

if st.button("🚀 Pasul 1: Generează Traducerea", type="primary"):
    if not uploaded_file:
        st.error("Încarcă un fișier Excel!")
    else:
        with st.spinner("Se traduce fișierul..."):
            wb, stats, new_terms = run_translation_pipeline(
                uploaded_file.getvalue(),
                master_dict
            )
            st.session_state.translated_wb = wb
            st.session_state.stats = stats
            
            df_data = []
            for src, trg in new_terms.items():
                df_data.append({"Validează": True, "Termen Română": src, "Traducere Engleză (Editabil)": trg})
            
            st.session_state.new_terms_df = pd.DataFrame(df_data)

if st.session_state.translated_wb is not None:
    stats = st.session_state.stats
    st.success(f"Statistici: Dicționar (MATCH): {stats['MATCH']} | Google Translate (MISSING): {stats['MISSING']} | Netraduse (NOT FOUND): {stats['NOT FOUND']}")

    st.markdown("### 📝 Pasul 2: Revizuire & Salvare Directă în Dicționar")

    if not st.session_state.new_terms_df.empty:
        edited_df = st.data_editor(
            st.session_state.new_terms_df,
            column_config={
                "Validează": st.column_config.CheckboxColumn("Aprobă pentru dicționar?", default=True),
                "Termen Română": st.column_config.TextColumn("Termen Română", disabled=True),
                "Traducere Engleză (Editabil)": st.column_config.TextColumn("Traducere Engleză (Editează dacă e cazul)")
            },
            hide_index=True,
            use_container_width=True
        )
    else:
        st.write("Toți termenii au existat deja în dicționar.")
        edited_df = pd.DataFrame()

    st.markdown("### 💾 Pasul 3: Salvează în Dicționar și Descarcă Excelul Tradus")
    
    if st.button("✅ Confirmă Termenii și Pregătește Descărcarea"):
        if not edited_df.empty:
            approved_dict = {}
            for _, row in edited_df.iterrows():
                if row["Validează"]:
                    approved_dict[row["Termen Română"]] = row["Traducere Engleză (Editabil)"]
            
            if approved_dict:
                with st.spinner("Se salvează termenii direct în dicționarul de pe GitHub..."):
                    if update_github_dictionary(approved_dict):
                        st.success(f"✅ {len(approved_dict)} termeni noi au fost salvați direct și permanent în dicționar!")

        buf = io.BytesIO()
        st.session_state.translated_wb.save(buf)
        buf.seek(0)

        st.download_button(
            label="📥 Descarcă Fișierul Excel Tradus (.xlsx)",
            data=buf,
            file_name=f"translated_{uploaded_file.name}",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
